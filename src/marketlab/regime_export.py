"""Workbook exports for saved market-regime research snapshots."""

from __future__ import annotations

import csv
import re
from contextlib import suppress
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Any
from zipfile import BadZipFile

from loguru import logger
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.table import Table, TableStyleInfo

from marketlab.prices import DEFAULT_PRICE_OBSERVATIONS_ROOT, load_price_history
from marketlab.regime import (
    ACTION_EXTENSION_ARMED_MAX_AGE_DAYS,
    DEFAULT_REGIME_SNAPSHOT_ROOT,
    EXTENSION_ARMED_SETUP_THRESHOLD,
    PRESSURE_FEATURE_SPECS,
    PRESSURE_MODULE_SPECS,
    REGIME_SNAPSHOT_FILENAME,
    RegimeError,
    RegimePostureState,
    RegimeSnapshot,
    derive_regime_postures,
    load_regime_snapshot,
)
from marketlab.snapshots.models import SignalTime, signal_time_value, snapshot_source_dates
from marketlab.versions import REGIME_VERSION

DEFAULT_REGIME_EXPORT_ROOT = Path("workbooks")
DEFAULT_REGIME_EXPORT_WORKBOOK_PATH = DEFAULT_REGIME_EXPORT_ROOT / f"regime-{REGIME_VERSION}.xlsx"
DEFAULT_REGIME_EXPORT_CSV_PATH = DEFAULT_REGIME_EXPORT_ROOT / f"regime-{REGIME_VERSION}.csv"
DEFAULT_REGIME_EXPORT_SHEET_NAME = "regime_data"
DEFAULT_REGIME_EXPORT_TABLE_NAME = "RegimeData"
ARMED_EXTENSION_MAX_AGE_DAYS = ACTION_EXTENSION_ARMED_MAX_AGE_DAYS

_ACTION_FEATURE_ORDER = (
    "short_momentum",
    "vix_turn",
    "vvix_turn",
    "term_structure_turn",
    "vvix_vs_vix",
    "volatility_normalization_confirmation",
    "participation_turn",
    "price_breadth_divergence",
    "hyg_tlt_turn",
    "credit_spread_turn",
    "equity_credit_divergence",
    "breadth_credit_confirmation",
    "rebound_not_confirmed_by_credit",
    "financial_conditions_turn",
)

_ACTION_MODULE_ORDER = (
    "trend_turn",
    "volatility_turn",
    "breadth_turn",
    "credit_turn",
    "liquidity_turn",
)

_BASE_COLUMNS = (
    "date",
    "price_date",
    "regime",
    "regime_score",
    "regime_confidence",
    "action_state",
    "action_path",
    "action_confidence",
    "action_reason",
    "event_volatility_context",
    "posture_state",
    "protected_since",
    "protect_path",
    "protect_reason",
    "redeploy_since",
    "posture_age",
    "max_drawdown_since_posture",
    "posture_reset_reason",
    "stretch_pressure",
    "washout_pressure",
    "extension_setup",
    "extension_elevated_since",
    "extension_elevated_age",
    "extension_armed",
    "extension_armed_since",
    "extension_armed_age",
    "extension_armed_decay",
    "extension_armed_source_strength",
    "fragility_context",
    "washout_setup",
    "protect_strength",
    "redeploy_strength",
    "net_action_score",
    "deterioration_score",
    "stabilization_score",
    "SPY_close",
    "QQQ_close",
    "SPY_indexed_100",
    "QQQ_indexed_100",
)

_MARKET_INPUT_COLUMNS = (
    "spy_return_20d",
    "spy_return_60d",
    "spy_ma50_distance",
    "spy_ma200_distance",
    "qqq_return_20d",
    "qqq_ma50_distance",
    "qqq_ma200_distance",
    "spy_realized_vol_20d",
    "vix_close",
    "vvix_close",
    "vix9d_vix_ratio",
    "vix_vix3m_ratio",
    "rsp_spy_relative_20d",
    "iwm_spy_relative_20d",
    "sp500_above_50d",
    "sp500_above_200d",
    "hyg_tlt_relative_20d",
)

_DATE_COLUMNS = frozenset(
    {
        "date",
        "price_date",
        "protected_since",
        "redeploy_since",
        "extension_elevated_since",
        "extension_armed_since",
    }
)
_PERCENT_COLUMNS = frozenset(
    {
        "spy_return_20d",
        "spy_return_60d",
        "spy_ma50_distance",
        "spy_ma200_distance",
        "qqq_return_20d",
        "qqq_ma50_distance",
        "qqq_ma200_distance",
        "spy_realized_vol_20d",
        "rsp_spy_relative_20d",
        "iwm_spy_relative_20d",
        "hyg_tlt_relative_20d",
    }
)


class RegimeExportError(ValueError):
    """Raised when regime exports cannot be produced."""


@dataclass(frozen=True, slots=True)
class RegimeExportResult:
    """Metadata for one regime workbook export."""

    start_date: date
    end_date: date
    workbook_path: Path
    csv_path: Path
    row_count: int
    column_count: int


def regime_export_paths(
    *,
    version: str | None = None,
    output_root: Path = DEFAULT_REGIME_EXPORT_ROOT,
) -> tuple[Path, Path]:
    """Return standard versioned XLSX and CSV export paths."""

    version_text = version or REGIME_VERSION
    return (
        output_root / f"regime-{_artifact_part(version_text)}.xlsx",
        output_root / f"regime-{_artifact_part(version_text)}.csv",
    )


def regime_outcome_report_path(
    *,
    outcome_group: str,
    version: str | None = None,
    output_root: Path = DEFAULT_REGIME_EXPORT_ROOT,
) -> Path:
    """Return the standard versioned text report path for one outcome group."""

    version_text = version or REGIME_VERSION
    return output_root / (
        f"outcomes-{_artifact_part(outcome_group)}-{_artifact_part(version_text)}.txt"
    )


def export_regime_workbook(
    *,
    start_date: date,
    end_date: date,
    workbook_path: Path = DEFAULT_REGIME_EXPORT_WORKBOOK_PATH,
    csv_path: Path = DEFAULT_REGIME_EXPORT_CSV_PATH,
    snapshot_root: Path = DEFAULT_REGIME_SNAPSHOT_ROOT,
    price_root: Path = DEFAULT_PRICE_OBSERVATIONS_ROOT,
    signal_time: SignalTime | str = SignalTime.EOD,
    snapshot_version: str | None = None,
) -> RegimeExportResult:
    """Export saved regime snapshots into refreshable XLSX and CSV files."""

    if start_date > end_date:
        raise RegimeExportError("Export start date must be on or before end date")

    snapshots = _load_snapshots(
        start_date=start_date,
        end_date=end_date,
        snapshot_root=snapshot_root,
        signal_time=signal_time,
        snapshot_version=snapshot_version,
    )
    if not snapshots:
        raise RegimeExportError(
            "No market regime snapshots found for "
            f"{start_date.isoformat()} through {end_date.isoformat()}"
    )

    headers = _headers_for_snapshots(snapshots)
    rows = _rows_for_snapshots(snapshots, price_root=price_root, headers=headers)
    workbook_temp_path = _temporary_export_path(workbook_path)
    csv_temp_path = _temporary_export_path(csv_path)
    try:
        _write_workbook(
            workbook_path=workbook_path,
            output_path=workbook_temp_path,
            sheet_name=DEFAULT_REGIME_EXPORT_SHEET_NAME,
            table_name=DEFAULT_REGIME_EXPORT_TABLE_NAME,
            headers=headers,
            rows=rows,
        )
        _write_csv(csv_path=csv_temp_path, headers=headers, rows=rows)
        _replace_export_file(temp_path=workbook_temp_path, final_path=workbook_path)
        _replace_export_file(temp_path=csv_temp_path, final_path=csv_path)
    except RegimeExportError:
        _cleanup_temp_export(workbook_temp_path)
        _cleanup_temp_export(csv_temp_path)
        raise

    logger.bind(
        start_date=start_date.isoformat(),
        end_date=end_date.isoformat(),
        workbook_path=str(workbook_path),
        csv_path=str(csv_path),
        row_count=len(rows),
        column_count=len(headers),
    ).success("Exported regime snapshots to workbook files")
    return RegimeExportResult(
        start_date=start_date,
        end_date=end_date,
        workbook_path=workbook_path,
        csv_path=csv_path,
        row_count=len(rows),
        column_count=len(headers),
    )


def _load_snapshots(
    *,
    start_date: date,
    end_date: date,
    snapshot_root: Path,
    signal_time: SignalTime | str,
    snapshot_version: str | None,
) -> tuple[RegimeSnapshot, ...]:
    signal_time_text = signal_time_value(signal_time)
    snapshots: list[RegimeSnapshot] = []
    for source_date in snapshot_source_dates(
        snapshot_root=snapshot_root,
        legacy_suffix=".json",
        filename=REGIME_SNAPSHOT_FILENAME,
    ):
        if not start_date <= source_date <= end_date:
            continue
        try:
            snapshot = load_regime_snapshot(
                source_date=source_date,
                snapshot_root=snapshot_root,
                signal_time=signal_time_text,
                snapshot_version=snapshot_version,
            )
        except RegimeError as exc:
            raise RegimeExportError(str(exc)) from exc
        if snapshot is not None:
            snapshots.append(snapshot)
    return tuple(sorted(snapshots, key=lambda snapshot: snapshot.source_date))


def _headers_for_snapshots(snapshots: tuple[RegimeSnapshot, ...]) -> tuple[str, ...]:
    pressure_feature_names = _ordered_names(
        known_names=tuple(spec.name for spec in PRESSURE_FEATURE_SPECS),
        seen_names=(
            feature.name
            for snapshot in snapshots
            for feature in snapshot.pressure_features
        ),
    )
    pressure_module_names = _ordered_names(
        known_names=tuple(spec.name for spec in PRESSURE_MODULE_SPECS),
        seen_names=(
            module.name
            for snapshot in snapshots
            for module in snapshot.pressure_modules
        ),
    )
    action_feature_names = _ordered_names(
        known_names=_ACTION_FEATURE_ORDER,
        seen_names=(
            feature.name
            for snapshot in snapshots
            for feature in snapshot.action_features
        ),
    )
    action_module_names = _ordered_names(
        known_names=_ACTION_MODULE_ORDER,
        seen_names=(
            module.name
            for snapshot in snapshots
            for module in snapshot.action_modules
        ),
    )

    headers: list[str] = [*_BASE_COLUMNS, *_MARKET_INPUT_COLUMNS]
    for name in pressure_feature_names:
        prefix = f"pressure__{_column_part(name)}"
        headers.extend(
            (
                f"{prefix}__raw",
                f"{prefix}__pctile",
                f"{prefix}__stretch",
                f"{prefix}__washout",
                f"{prefix}__samples",
            )
        )
    for name in pressure_module_names:
        prefix = f"pressure_module__{_column_part(name)}"
        headers.extend(
            (
                f"{prefix}__stretch",
                f"{prefix}__washout",
                f"{prefix}__feature_count",
                f"{prefix}__reason",
            )
        )
    for name in action_feature_names:
        prefix = f"action__{_column_part(name)}"
        headers.extend(
            (
                f"{prefix}__protect",
                f"{prefix}__redeploy",
                f"{prefix}__kind",
                f"{prefix}__reason",
            )
        )
    for name in action_module_names:
        prefix = f"action_module__{_column_part(name)}"
        headers.extend(
            (
                f"{prefix}__protect",
                f"{prefix}__redeploy",
                f"{prefix}__feature_count",
                f"{prefix}__reason",
            )
        )
    return tuple(headers)


def _rows_for_snapshots(
    snapshots: tuple[RegimeSnapshot, ...],
    *,
    price_root: Path,
    headers: tuple[str, ...],
) -> list[dict[str, Any]]:
    spy_prices = _prices_by_date(symbol="SPY", price_root=price_root)
    qqq_prices = _prices_by_date(symbol="QQQ", price_root=price_root)
    spy_base = _first_snapshot_price(snapshots, prices_by_date=spy_prices)
    qqq_base = _first_snapshot_price(snapshots, prices_by_date=qqq_prices)
    postures = {
        posture.source_date: posture
        for posture in derive_regime_postures(
            snapshots=snapshots,
            spy_history=tuple(load_price_history(symbol="SPY", price_root=price_root)),
        )
    }
    armed_rows = _armed_extension_memory_rows(snapshots)

    rows: list[dict[str, Any]] = []
    for snapshot in snapshots:
        row = _base_row(
            snapshot,
            spy_prices=spy_prices,
            qqq_prices=qqq_prices,
            spy_base=spy_base,
            qqq_base=qqq_base,
        )
        _add_posture(row, postures.get(snapshot.source_date))
        row.update(armed_rows.get(snapshot.source_date, {}))
        _add_pressure_features(row, snapshot)
        _add_pressure_modules(row, snapshot)
        _add_action_features(row, snapshot)
        _add_action_modules(row, snapshot)
        rows.append({header: row.get(header) for header in headers})
    return rows


def _add_posture(row: dict[str, Any], posture: RegimePostureState | None) -> None:
    if posture is None:
        return
    row.update(
        {
            "posture_state": posture.posture_state,
            "protected_since": posture.protected_since,
            "protect_path": posture.protect_path,
            "protect_reason": posture.protect_reason,
            "redeploy_since": posture.redeploy_since,
            "posture_age": posture.posture_age,
            "max_drawdown_since_posture": posture.max_drawdown_since_posture,
            "posture_reset_reason": posture.reset_reason,
        }
    )


def _armed_extension_memory_rows(
    snapshots: tuple[RegimeSnapshot, ...],
) -> dict[date, dict[str, Any]]:
    elevated_since: date | None = None
    elevated_index: int | None = None
    armed_since: date | None = None
    armed_index: int | None = None
    source_strength: float | None = None
    rows: dict[date, dict[str, Any]] = {}
    for index, snapshot in enumerate(snapshots):
        setup_score = snapshot.extension_setup_score
        if _armed_extension_setup_active(setup_score):
            if elevated_since is None:
                elevated_since = snapshot.source_date
                elevated_index = index
                armed_since = snapshot.source_date
                armed_index = index
                source_strength = setup_score
        else:
            elevated_since = None
            elevated_index = None
            if armed_index is not None and index - armed_index > ARMED_EXTENSION_MAX_AGE_DAYS:
                armed_since = None
                armed_index = None
                source_strength = None

        age = index - armed_index if armed_index is not None else None
        elevated_age = index - elevated_index if elevated_index is not None else None
        is_armed = age is not None and age <= ARMED_EXTENSION_MAX_AGE_DAYS
        decay = (
            max(0.0, 1.0 - age / ARMED_EXTENSION_MAX_AGE_DAYS)
            if is_armed and age is not None
            else None
        )
        rows[snapshot.source_date] = {
            "extension_elevated_since": elevated_since,
            "extension_elevated_age": elevated_age,
            "extension_armed": is_armed,
            "extension_armed_since": armed_since,
            "extension_armed_age": age,
            "extension_armed_decay": decay,
            "extension_armed_source_strength": source_strength,
        }
    return rows


def _armed_extension_setup_active(setup_score: float | None) -> bool:
    return setup_score is not None and setup_score >= EXTENSION_ARMED_SETUP_THRESHOLD


def _base_row(
    snapshot: RegimeSnapshot,
    *,
    spy_prices: dict[date, float],
    qqq_prices: dict[date, float],
    spy_base: float | None,
    qqq_base: float | None,
) -> dict[str, Any]:
    spy_close = spy_prices.get(snapshot.price_date) if snapshot.price_date is not None else None
    qqq_close = qqq_prices.get(snapshot.price_date) if snapshot.price_date is not None else None
    return {
        "date": snapshot.source_date,
        "price_date": snapshot.price_date,
        "regime": snapshot.regime,
        "regime_score": snapshot.score,
        "regime_confidence": snapshot.confidence,
        "action_state": snapshot.action_state,
        "action_path": snapshot.action_path,
        "action_confidence": snapshot.action_confidence,
        "action_reason": snapshot.action_reason,
        "event_volatility_context": snapshot.event_volatility_context,
        "stretch_pressure": snapshot.stretch_pressure_score,
        "washout_pressure": snapshot.washout_pressure_score,
        "extension_setup": snapshot.extension_setup_score,
        "fragility_context": snapshot.fragility_context_score,
        "washout_setup": snapshot.washout_setup_score,
        "protect_strength": snapshot.protect_strength,
        "redeploy_strength": snapshot.redeploy_strength,
        "net_action_score": snapshot.action_score,
        "deterioration_score": snapshot.deterioration_score,
        "stabilization_score": snapshot.stabilization_score,
        "SPY_close": spy_close,
        "QQQ_close": qqq_close,
        "SPY_indexed_100": _indexed_value(spy_close, spy_base),
        "QQQ_indexed_100": _indexed_value(qqq_close, qqq_base),
        "spy_return_20d": snapshot.spy_return_20d,
        "spy_return_60d": snapshot.spy_return_60d,
        "spy_ma50_distance": snapshot.spy_ma50_distance,
        "spy_ma200_distance": snapshot.spy_ma200_distance,
        "qqq_return_20d": snapshot.qqq_return_20d,
        "qqq_ma50_distance": snapshot.qqq_ma50_distance,
        "qqq_ma200_distance": snapshot.qqq_ma200_distance,
        "spy_realized_vol_20d": snapshot.spy_realized_vol_20d,
        "vix_close": snapshot.vix_close,
        "vvix_close": snapshot.vvix_close,
        "vix9d_vix_ratio": snapshot.vix9d_vix_ratio,
        "vix_vix3m_ratio": snapshot.vix_vix3m_ratio,
        "rsp_spy_relative_20d": snapshot.rsp_spy_relative_20d,
        "iwm_spy_relative_20d": snapshot.iwm_spy_relative_20d,
        "sp500_above_50d": snapshot.sp500_above_50d,
        "sp500_above_200d": snapshot.sp500_above_200d,
        "hyg_tlt_relative_20d": snapshot.hyg_tlt_relative_20d,
    }


def _add_pressure_features(row: dict[str, Any], snapshot: RegimeSnapshot) -> None:
    for feature in snapshot.pressure_features:
        prefix = f"pressure__{_column_part(feature.name)}"
        row[f"{prefix}__raw"] = feature.raw_value
        row[f"{prefix}__pctile"] = feature.percentile
        row[f"{prefix}__stretch"] = feature.stretch_score
        row[f"{prefix}__washout"] = feature.washout_score
        row[f"{prefix}__samples"] = feature.sample_count


def _add_pressure_modules(row: dict[str, Any], snapshot: RegimeSnapshot) -> None:
    for module in snapshot.pressure_modules:
        prefix = f"pressure_module__{_column_part(module.name)}"
        row[f"{prefix}__stretch"] = module.stretch_score
        row[f"{prefix}__washout"] = module.washout_score
        row[f"{prefix}__feature_count"] = module.feature_count
        row[f"{prefix}__reason"] = module.reason


def _add_action_features(row: dict[str, Any], snapshot: RegimeSnapshot) -> None:
    for feature in snapshot.action_features:
        prefix = f"action__{_column_part(feature.name)}"
        row[f"{prefix}__protect"] = feature.deterioration_score
        row[f"{prefix}__redeploy"] = feature.stabilization_score
        row[f"{prefix}__kind"] = feature.evidence_kind
        row[f"{prefix}__reason"] = feature.reason


def _add_action_modules(row: dict[str, Any], snapshot: RegimeSnapshot) -> None:
    for module in snapshot.action_modules:
        prefix = f"action_module__{_column_part(module.name)}"
        row[f"{prefix}__protect"] = module.deterioration_score
        row[f"{prefix}__redeploy"] = module.stabilization_score
        row[f"{prefix}__feature_count"] = module.feature_count
        row[f"{prefix}__reason"] = module.reason


def _write_csv(
    *,
    csv_path: Path,
    headers: tuple[str, ...],
    rows: list[dict[str, Any]],
) -> None:
    try:
        csv_path.parent.mkdir(parents=True, exist_ok=True)
        with csv_path.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=headers, extrasaction="ignore")
            writer.writeheader()
            for row in rows:
                writer.writerow({header: _csv_value(row.get(header)) for header in headers})
    except OSError as exc:
        raise RegimeExportError(f"Failed to write CSV export: {csv_path}") from exc


def _write_workbook(
    *,
    workbook_path: Path,
    output_path: Path,
    sheet_name: str,
    table_name: str,
    headers: tuple[str, ...],
    rows: list[dict[str, Any]],
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    if workbook_path.is_file():
        try:
            workbook = load_workbook(workbook_path)
        except (OSError, InvalidFileException, BadZipFile) as exc:
            raise RegimeExportError(f"Failed to load workbook export: {workbook_path}") from exc
        worksheet = (
            workbook[sheet_name]
            if sheet_name in workbook.sheetnames
            else workbook.create_sheet(sheet_name, 0)
        )
    else:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = sheet_name

    _clear_sheet_values(worksheet, min_rows=len(rows) + 1, min_cols=len(headers))
    for column_index, header in enumerate(headers, start=1):
        cell = worksheet.cell(row=1, column=column_index, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    for row_index, row in enumerate(rows, start=2):
        for column_index, header in enumerate(headers, start=1):
            cell = worksheet.cell(row=row_index, column=column_index, value=row.get(header))
            number_format = _number_format(header)
            if number_format is not None:
                cell.number_format = number_format

    table_ref = _range_ref(row_count=len(rows) + 1, column_count=len(headers))
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = table_ref
    _replace_table(worksheet, table_name=table_name, table_ref=table_ref)
    _size_columns(worksheet, headers=headers, rows=rows)
    try:
        workbook.save(output_path)
    except (OSError, InvalidFileException, BadZipFile) as exc:
        raise RegimeExportError(f"Failed to write workbook export: {output_path}") from exc


def _temporary_export_path(path: Path) -> Path:
    return path.with_name(f".{path.name}.tmp")


def _replace_export_file(*, temp_path: Path, final_path: Path) -> None:
    try:
        final_path.parent.mkdir(parents=True, exist_ok=True)
        temp_path.replace(final_path)
    except OSError as exc:
        raise RegimeExportError(f"Failed to replace export file: {final_path}") from exc


def _cleanup_temp_export(path: Path) -> None:
    with suppress(OSError):
        path.unlink()


def _clear_sheet_values(worksheet: Any, *, min_rows: int, min_cols: int) -> None:
    max_row = max(worksheet.max_row, min_rows)
    max_column = max(worksheet.max_column, min_cols)
    for row in worksheet.iter_rows(
        min_row=1,
        max_row=max_row,
        min_col=1,
        max_col=max_column,
    ):
        for cell in row:
            cell.value = None


def _replace_table(worksheet: Any, *, table_name: str, table_ref: str) -> None:
    if table_name in worksheet.tables:
        del worksheet.tables[table_name]
    table = Table(displayName=table_name, ref=table_ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    worksheet.add_table(table)


def _size_columns(
    worksheet: Any,
    *,
    headers: tuple[str, ...],
    rows: list[dict[str, Any]],
) -> None:
    for column_index, header in enumerate(headers, start=1):
        values = [header]
        values.extend(_column_width_value(row.get(header)) for row in rows[:100])
        width = min(max(len(value) for value in values) + 2, 42)
        worksheet.column_dimensions[get_column_letter(column_index)].width = width


def _column_width_value(value: Any) -> str:
    return "" if value is None else str(value)


def _range_ref(*, row_count: int, column_count: int) -> str:
    return f"A1:{get_column_letter(column_count)}{row_count}"


def _prices_by_date(*, symbol: str, price_root: Path) -> dict[date, float]:
    return {
        row.observed_date: row.return_close
        for row in load_price_history(symbol=symbol, price_root=price_root)
    }


def _first_snapshot_price(
    snapshots: tuple[RegimeSnapshot, ...],
    *,
    prices_by_date: dict[date, float],
) -> float | None:
    for snapshot in snapshots:
        if snapshot.price_date is None:
            continue
        price = prices_by_date.get(snapshot.price_date)
        if price not in (None, 0):
            return price
    return None


def _indexed_value(close: float | None, base: float | None) -> float | None:
    if close is None or base in (None, 0):
        return None
    return close / base * 100.0


def _csv_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, date):
        return value.isoformat()
    return value


def _ordered_names(
    *,
    known_names: tuple[str, ...],
    seen_names: Any,
) -> tuple[str, ...]:
    seen = tuple(dict.fromkeys(str(name) for name in seen_names))
    seen_set = set(seen)
    known = tuple(name for name in known_names if name in seen_set)
    unknown = tuple(sorted(name for name in seen if name not in set(known_names)))
    return known + unknown


def _column_part(value: str) -> str:
    return re.sub(r"_+", "_", re.sub(r"[^0-9A-Za-z]+", "_", value.strip()).strip("_")).lower()


def _artifact_part(value: str) -> str:
    text = str(value).strip()
    safe = re.sub(r"_+", "_", re.sub(r"[^0-9A-Za-z._-]+", "_", text).strip("_"))
    if not safe or safe in {".", ".."}:
        raise RegimeExportError(f"Invalid artifact path component: {value!r}")
    return safe


def _number_format(header: str) -> str | None:
    if header in _DATE_COLUMNS:
        return "yyyy-mm-dd"
    if header in _PERCENT_COLUMNS or header.endswith("__pctile"):
        return "0.0%"
    if (
        header.endswith("__stretch")
        or header.endswith("__washout")
        or header.endswith("__protect")
        or header.endswith("__redeploy")
        or header.endswith("_score")
        or header.endswith("_strength")
        or header in {"SPY_indexed_100", "QQQ_indexed_100"}
    ):
        return "0.00"
    if header.endswith("_close") or header.endswith("_ratio"):
        return "0.00"
    return None
