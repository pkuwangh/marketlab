import csv
from datetime import date, datetime, timedelta
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from marketlab.prices import DailyPriceObservation, write_price_observations
from marketlab.regime import (
    RegimeActionFeature,
    RegimeActionModule,
    RegimeComponentScore,
    RegimePressureFeature,
    RegimePressureModule,
    RegimeSnapshot,
    write_regime_snapshot,
)
from marketlab.regime_export import (
    RegimeExportError,
    _column_width_value,
    export_regime_workbook,
    regime_export_paths,
    regime_outcome_report_path,
)


def test_export_regime_workbook_writes_refreshable_data_sheet_and_csv(
    tmp_path: Path,
) -> None:
    snapshot_root = tmp_path / "snapshots"
    price_root = tmp_path / "prices"
    workbook_path = tmp_path / "regime.xlsx"
    csv_path = tmp_path / "regime.csv"
    first_date = date(2026, 9, 4)
    second_date = date(2026, 9, 8)
    _write_prices(price_root, symbol="SPY", start=first_date, closes=(100.0, 103.0))
    _write_prices(price_root, symbol="QQQ", start=first_date, closes=(200.0, 202.0))
    write_regime_snapshot(
        _snapshot(
            source_date=first_date,
            spy_return_20d=0.01,
            extension_setup_score=0.76,
            washout_setup_score=0.80,
        ),
        snapshot_root=snapshot_root,
    )
    write_regime_snapshot(
        _snapshot(
            source_date=second_date,
            spy_return_20d=0.03,
            extension_setup_score=0.40,
            washout_setup_score=0.40,
        ),
        snapshot_root=snapshot_root,
    )

    result = export_regime_workbook(
        start_date=first_date,
        end_date=second_date,
        workbook_path=workbook_path,
        csv_path=csv_path,
        snapshot_root=snapshot_root,
        price_root=price_root,
    )

    assert result.row_count == 2
    assert result.column_count > 40
    assert result.workbook_path == workbook_path
    assert result.csv_path == csv_path
    workbook = load_workbook(workbook_path)
    worksheet = workbook["regime_data"]
    headers = [cell.value for cell in worksheet[1]]
    first_row = dict(zip(headers, (cell.value for cell in worksheet[2]), strict=True))
    assert _as_date(first_row["date"]) == first_date
    assert first_row["SPY_close"] == 100.0
    assert first_row["SPY_indexed_100"] == 100.0
    assert first_row["pressure__vix_close__pctile"] == 0.20
    assert first_row["pressure_module__vol_complacency__stretch"] == 0.80
    assert first_row["action__vvix_turn__protect"] == 0.70
    assert first_row["action_module__volatility_turn__redeploy"] == 0.10
    assert first_row["action_path"] == "extension_unwind"
    assert first_row["extension_setup"] == 0.76
    assert first_row["fragility_context"] == 0.20
    assert first_row["washout_setup"] == 0.80
    assert first_row["event_volatility_context"] == "near_term_event_premium"
    assert first_row["posture_state"] == "protected"
    assert _as_date(first_row["protected_since"]) == first_date
    assert first_row["protect_path"] == "extension_unwind"
    assert first_row["posture_age"] == 0
    assert first_row["max_drawdown_since_posture"] == 0.0
    assert _as_date(first_row["extension_elevated_since"]) == first_date
    assert first_row["extension_elevated_age"] == 0
    assert first_row["extension_armed"] is True
    assert _as_date(first_row["extension_armed_since"]) == first_date
    assert first_row["extension_armed_age"] == 0
    assert first_row["extension_armed_decay"] == 1.0
    assert first_row["extension_armed_source_strength"] == 0.76
    assert first_row["extension_arm_rule"] == "strong"
    assert first_row["extension_arm_hit_count"] == 1
    assert first_row["extension_arm_window"] == 1
    assert first_row["washout_armed"] is True
    assert _as_date(first_row["washout_armed_since"]) == first_date
    assert first_row["washout_armed_age"] == 0
    assert first_row["washout_armed_decay"] == 1.0
    assert first_row["washout_armed_source_strength"] == 0.80
    assert first_row["washout_arm_rule"] == "strong"
    assert first_row["washout_memory_score"] == 0.80
    second_row = dict(zip(headers, (cell.value for cell in worksheet[3]), strict=True))
    assert second_row["extension_elevated_since"] is None
    assert second_row["extension_elevated_age"] is None
    assert second_row["extension_armed"] is True
    assert _as_date(second_row["extension_armed_since"]) == first_date
    assert second_row["extension_armed_age"] == 1
    assert second_row["extension_arm_rule"] == "strong"
    assert second_row["washout_armed"] is True
    assert _as_date(second_row["washout_armed_since"]) == first_date
    assert second_row["washout_armed_age"] == 1
    assert second_row["washout_memory_score"] > 0.40
    last_column = worksheet.cell(row=1, column=len(headers)).column_letter
    assert worksheet.tables["RegimeData"].ref == f"A1:{last_column}3"
    with csv_path.open(encoding="utf-8", newline="") as handle:
        csv_rows = list(csv.DictReader(handle))
    assert csv_rows[0]["date"] == first_date.isoformat()
    assert csv_rows[0]["SPY_close"] == "100.0"
    assert csv_rows[0]["event_volatility_context"] == "near_term_event_premium"
    assert csv_rows[0]["posture_state"] == "protected"
    assert csv_rows[0]["protect_path"] == "extension_unwind"
    assert csv_rows[0]["action_path"] == "extension_unwind"
    assert csv_rows[0]["extension_elevated_age"] == "0"
    assert csv_rows[1]["extension_elevated_age"] == ""
    assert csv_rows[0]["extension_armed"] == "True"
    assert csv_rows[1]["extension_armed_age"] == "1"
    assert csv_rows[0]["extension_arm_rule"] == "strong"
    assert csv_rows[0]["pressure__vix_close__pctile"] == "0.2"


def test_export_regime_workbook_does_not_rearm_continuous_extension_run(
    tmp_path: Path,
) -> None:
    snapshot_root = tmp_path / "snapshots"
    price_root = tmp_path / "prices"
    workbook_path = tmp_path / "regime.xlsx"
    csv_path = tmp_path / "regime.csv"
    first_date = date(2026, 1, 2)
    row_count = 44
    dates = tuple(first_date + timedelta(days=index * 4) for index in range(row_count))
    _write_prices(
        price_root,
        symbol="SPY",
        start=first_date,
        closes=tuple(100.0 + index for index in range(row_count)),
    )
    _write_prices(
        price_root,
        symbol="QQQ",
        start=first_date,
        closes=tuple(200.0 + index for index in range(row_count)),
    )
    for source_date in dates:
        write_regime_snapshot(
            _snapshot(source_date=source_date, extension_setup_score=0.65),
            snapshot_root=snapshot_root,
        )

    export_regime_workbook(
        start_date=dates[0],
        end_date=dates[-1],
        workbook_path=workbook_path,
        csv_path=csv_path,
        snapshot_root=snapshot_root,
        price_root=price_root,
    )

    workbook = load_workbook(workbook_path)
    worksheet = workbook["regime_data"]
    headers = [cell.value for cell in worksheet[1]]
    last_row = dict(zip(headers, (cell.value for cell in worksheet[row_count + 1]), strict=True))
    assert _as_date(last_row["extension_elevated_since"]) == first_date
    assert last_row["extension_elevated_age"] == row_count - 1
    assert last_row["extension_armed"] is False
    assert _as_date(last_row["extension_armed_since"]) == dates[2]
    assert last_row["extension_armed_age"] == row_count - 3
    assert last_row["extension_armed_decay"] is None
    assert last_row["extension_arm_rule"] == "persistent"


def test_export_regime_workbook_preserves_other_sheets_and_clears_stale_rows(
    tmp_path: Path,
) -> None:
    snapshot_root = tmp_path / "snapshots"
    price_root = tmp_path / "prices"
    workbook_path = tmp_path / "regime.xlsx"
    csv_path = tmp_path / "regime.csv"
    first_date = date(2026, 9, 4)
    second_date = date(2026, 9, 8)
    _write_prices(price_root, symbol="SPY", start=first_date, closes=(100.0, 103.0))
    _write_prices(price_root, symbol="QQQ", start=first_date, closes=(200.0, 202.0))
    write_regime_snapshot(_snapshot(source_date=first_date), snapshot_root=snapshot_root)
    write_regime_snapshot(_snapshot(source_date=second_date), snapshot_root=snapshot_root)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "manual_charts"
    worksheet["A1"] = "keep this sheet"
    workbook.save(workbook_path)
    export_regime_workbook(
        start_date=first_date,
        end_date=second_date,
        workbook_path=workbook_path,
        csv_path=csv_path,
        snapshot_root=snapshot_root,
        price_root=price_root,
    )
    export_regime_workbook(
        start_date=first_date,
        end_date=first_date,
        workbook_path=workbook_path,
        csv_path=csv_path,
        snapshot_root=snapshot_root,
        price_root=price_root,
    )

    refreshed = load_workbook(workbook_path)
    assert refreshed["manual_charts"]["A1"].value == "keep this sheet"
    data_sheet = refreshed["regime_data"]
    assert _as_date(data_sheet["A2"].value) == first_date
    assert data_sheet["A3"].value is None
    assert data_sheet.tables["RegimeData"].ref.endswith("2")
    with csv_path.open(encoding="utf-8", newline="") as handle:
        assert len(list(csv.DictReader(handle))) == 1


def test_export_regime_workbook_keeps_csv_unchanged_when_workbook_load_fails(
    tmp_path: Path,
) -> None:
    snapshot_root = tmp_path / "snapshots"
    price_root = tmp_path / "prices"
    workbook_path = tmp_path / "regime.xlsx"
    csv_path = tmp_path / "regime.csv"
    source_date = date(2026, 9, 4)
    _write_prices(price_root, symbol="SPY", start=source_date, closes=(100.0,))
    _write_prices(price_root, symbol="QQQ", start=source_date, closes=(200.0,))
    write_regime_snapshot(_snapshot(source_date=source_date), snapshot_root=snapshot_root)
    workbook_path.write_text("not a workbook", encoding="utf-8")
    csv_path.write_text("stale\n", encoding="utf-8")

    with pytest.raises(RegimeExportError, match="Failed to load workbook export"):
        export_regime_workbook(
            start_date=source_date,
            end_date=source_date,
            workbook_path=workbook_path,
            csv_path=csv_path,
            snapshot_root=snapshot_root,
            price_root=price_root,
        )

    assert csv_path.read_text(encoding="utf-8") == "stale\n"
    assert not (tmp_path / ".regime.xlsx.tmp").exists()
    assert not (tmp_path / ".regime.csv.tmp").exists()


def test_column_width_value_preserves_zero_values() -> None:
    assert _column_width_value(0) == "0"
    assert _column_width_value(0.0) == "0.0"
    assert _column_width_value(None) == ""


def test_versioned_regime_export_paths_use_artifact_version() -> None:
    assert regime_export_paths(version="v1.0") == (
        Path("workbooks/regime-v1.0.xlsx"),
        Path("workbooks/regime-v1.0.csv"),
    )
    assert regime_outcome_report_path(outcome_group="action", version="v1.0") == Path(
        "workbooks/outcomes-action-v1.0.txt"
    )


def _snapshot(
    *,
    source_date: date,
    spy_return_20d: float = 0.01,
    extension_setup_score: float = 0.50,
    washout_setup_score: float = 0.11,
    action_path: str | None = "extension_unwind",
) -> RegimeSnapshot:
    return RegimeSnapshot(
        source_date=source_date,
        feature_version="fixture_features",
        model_version="fixture_model",
        price_date=source_date,
        regime="risk_on",
        score=2,
        confidence="medium",
        components=(
            RegimeComponentScore(
                name="trend",
                score=2,
                status="constructive",
                reason="fixture trend",
            ),
        ),
        missing_symbols=(),
        spy_return_20d=spy_return_20d,
        spy_return_60d=0.08,
        spy_ma50_distance=0.02,
        spy_ma200_distance=0.09,
        spy_realized_vol_20d=0.12,
        qqq_return_20d=0.02,
        qqq_ma50_distance=0.03,
        qqq_ma200_distance=0.10,
        rsp_spy_relative_20d=-0.01,
        iwm_spy_relative_20d=-0.02,
        hyg_tlt_relative_20d=0.01,
        vix_close=15.0,
        vvix_close=84.0,
        vix9d_vix_ratio=0.94,
        vix_vix3m_ratio=0.78,
        sp500_above_50d=65.0,
        sp500_above_200d=70.0,
        action_state="scale_down_protect",
        action_path=action_path,
        action_confidence="medium",
        action_reason="extension_unwind protect: fixture action",
        event_volatility_context="near_term_event_premium",
        stretch_pressure_score=0.62,
        washout_pressure_score=0.11,
        extension_setup_score=extension_setup_score,
        fragility_context_score=0.20,
        washout_setup_score=washout_setup_score,
        deterioration_score=0.70,
        stabilization_score=0.10,
        protect_strength=0.70,
        redeploy_strength=0.10,
        action_score=-0.60,
        pressure_features=(
            RegimePressureFeature(
                source_date=source_date,
                pressure_version="fixture_pressure",
                family="volatility",
                name="vix_close",
                raw_value=15.0,
                percentile=0.20,
                normalized_score=0.80,
                stretch_score=0.80,
                washout_score=0.20,
                sample_count=126,
                reason="fixture pressure",
            ),
        ),
        pressure_modules=(
            RegimePressureModule(
                source_date=source_date,
                pressure_version="fixture_pressure",
                family="volatility",
                name="vol_complacency",
                stretch_score=0.80,
                washout_score=0.20,
                feature_count=1,
                feature_names=("vix_close",),
                reason="fixture pressure module",
            ),
        ),
        action_features=(
            RegimeActionFeature(
                family="volatility",
                name="vvix_turn",
                deterioration_score=0.70,
                stabilization_score=0.10,
                reason="fixture action feature",
            ),
        ),
        action_modules=(
            RegimeActionModule(
                family="volatility",
                name="volatility_turn",
                deterioration_score=0.70,
                stabilization_score=0.10,
                feature_count=1,
                feature_names=("vvix_turn",),
                reason="fixture action module",
            ),
        ),
    )


def _write_prices(
    price_root: Path,
    *,
    symbol: str,
    start: date,
    closes: tuple[float, ...],
) -> None:
    rows = tuple(
        DailyPriceObservation(
            source="fixture",
            source_symbol=symbol,
            symbol=symbol,
            observed_date=start + timedelta(days=index * 4),
            open=close,
            high=close,
            low=close,
            close=close,
            adjusted_close=close,
            volume=1_000_000,
        )
        for index, close in enumerate(closes)
    )
    write_price_observations(rows, output_root=price_root)


def _as_date(value: object) -> date:
    if isinstance(value, datetime):
        return value.date()
    assert isinstance(value, date)
    return value
